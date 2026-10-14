import json, sys, itertools
sys.path.insert(0, sys.argv[3])
from content import APPS, KEY, FLOWS, ENTERPRISE
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

rows = json.load(open(sys.argv[1])); out = sys.argv[2]
by = {r['tech']: r for r in rows}

DOM_RULES = [  # prefix -> domain (longest first)
 ('l10n','ภาษี/ประเทศ'),('account','บัญชี'),('analytic','บัญชี'),('payment','บัญชี'),('snailmail_account','บัญชี'),
 ('sales_team','ขาย'),('sale','ขาย'),('crm','CRM'),('purchase','จัดซื้อ'),('stock','คลังสินค้า'),('delivery','คลังสินค้า'),
 ('repair','คลังสินค้า'),('mrp','ผลิต'),('pos','POS'),('point_of_sale','POS'),('hr','บุคคล'),('lunch','บุคคล'),
 ('project','โครงการ'),('website','เว็บไซต์'),('im_livechat','เว็บไซต์'),('mass_mailing','การตลาด'),('sms','การตลาด'),
 ('marketing','การตลาด'),('link_tracker','การตลาด'),('utm','การตลาด'),('social','การตลาด'),('survey','การตลาด'),
 ('event','อีเว้นท์'),('fleet','รถ'),('maintenance','ซ่อมบำรุง'),('product','สินค้า'),('uom','สินค้า'),('loyalty','สินค้า'),
 ('gamification','บุคคล'),('resource','บุคคล'),
]
NONBIZ = {'เทคนิค','สินค้า','ภาษี/ประเทศ'}
def dom(t):
    for p,d in sorted(DOM_RULES, key=lambda x:-len(x[0])):
        if t==p or t.startswith(p+'_') or t.startswith(p): return d
    return 'เทคนิค'
apps = [r for r in rows if r['application']]
app_techs = {r['tech'] for r in apps}

def closure(t, seen=None):
    seen = set() if seen is None else seen
    for d in by.get(t,{}).get('depends',[]):
        if d not in seen and d in by: seen.add(d); closure(d, seen)
    return seen

pairs=[]
for r in rows:
    t=r['tech']; ds = {dom(t)} | {dom(d) for d in r['depends']}
    biz = sorted(ds - NONBIZ)
    r['domain']=dom(t); r['links']=biz
    if t.startswith('l10n_'): g='ภาษี/ประเทศ (l10n)'
    elif t.startswith('test'): g='โมดูลทดสอบ'
    elif r['application']: g='App หลัก'
    elif len(biz)>=2: g='ตัวเชื่อมระหว่างระบบ'
    elif r['domain']!='เทคนิค': g='ส่วนเสริม'
    else: g='เทคนิค/โครงสร้าง'
    r['group']=g
    if g=='ตัวเชื่อมระหว่างระบบ':
        for a,b in itertools.combinations(biz,2): pairs.append((t,a,b))
    r['req_apps']=sorted(a for a in closure(t) if a in app_techs)

# ---- styles
FN='Phetsarath OT'
GREEN='1E7B3C'; LIGHT='E8F5EC'; MID='C6E6CF'
f=lambda **k: Font(name=FN, size=k.pop('size',12), **k)
thin=Side(style='thin', color='9BBFA6'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
HFILL=PatternFill('solid', fgColor=GREEN); ZFILL=PatternFill('solid', fgColor=LIGHT); INFILL=PatternFill('solid', fgColor='FFF2CC')
WR=Alignment(wrap_text=True, vertical='top'); CT=Alignment(horizontal='center', vertical='top', wrap_text=True)

def title(ws, text, sub, span):
    ws['A1']=text; ws['A1'].font=f(size=16,bold=True,color=GREEN)
    ws['A2']=sub; ws['A2'].font=f(size=12,italic=True,color='555555'); ws['A2'].alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[2].height=36
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=span)
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=span)
    ws.row_dimensions[1].height=26

def header(ws, row, cols):
    for i,(h,w) in enumerate(cols,1):
        c=ws.cell(row=row,column=i,value=h); c.font=f(bold=True,color='FFFFFF'); c.fill=HFILL; c.alignment=CT; c.border=B
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[row].height=36

def body(ws, r0, data, center_cols=()):
    for i,rowv in enumerate(data):
        for j,v in enumerate(rowv,1):
            c=ws.cell(row=r0+i,column=j,value=v); c.font=f(); c.border=B
            c.alignment = CT if j in center_cols else WR
            if i%2: c.fill=ZFILL

def page(ws, hdr_row, landscape=True):
    ws.page_setup.orientation='landscape' if landscape else 'portrait'
    ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.print_title_rows=f'{hdr_row}:{hdr_row}'
    ws.page_margins.left=ws.page_margins.right=0.4; ws.page_margins.top=ws.page_margins.bottom=0.5
    ws.oddFooter.center.text='หน้า &P / &N'; ws.oddFooter.center.font=FN
    ws.freeze_panes=ws.cell(row=hdr_row+1,column=1)
    ws.sheet_view.showGridLines=False

wb=Workbook()
SRC='ที่มา: Odoo Community 20.0 (github.com/odoo/odoo branch 20.0, License LGPL-3) — นับจาก __manifest__.py, models/, wizard/, views/'

# ---- Sheet ทุก Module (built first so others can reference)
wa=wb.active; wa.title='ทุก Module'
cols=[('ลำดับ',7),('ชื่อเทคนิค',30),('ชื่อ',30),('กลุ่ม',20),('ระบบ',13),('เชื่อมระบบ',24),('ต้องมี App',30),('หมวด Odoo',24),('ทำอะไร',60),('Model ใหม่',10),('Model ที่ขยาย',10),('เมนู',8),('ต้องติดตั้งก่อน (depends)',40)]
title(wa,'Odoo Community 20.0 — รายการ Module ทั้งหมด',SRC,len(cols))
header(wa,4,cols)
data=[]
for i,r in enumerate(rows,1):
    what = KEY.get(r['tech']) or (APPS[r['tech']][0] if r['tech'] in APPS else (r['summary'] or r['name']))
    data.append([i,r['tech'],r['name'],r['group'],r['domain'],' + '.join(r['links']) if r['group']=='ตัวเชื่อมระหว่างระบบ' else '',
                 ', '.join(r['req_apps']),r['category'],what,len(r['new_models']),len(r['ext_models']),r['menus'],', '.join(r['depends'])])
body(wa,5,data,center_cols=(1,10,11,12))
N=len(rows); last=4+N
wa.auto_filter.ref=f'A4:{get_column_letter(len(cols))}{last}'
page(wa,4)

# ---- Sheet ตัวเชื่อม pairs
wp=wb.create_sheet('คู่เชื่อม (ข้อมูล)')
title(wp,'คู่ระบบที่ถูกเชื่อมด้วย Module ตัวเชื่อม','1 แถว = Module ตัวเชื่อม 1 ตัว ที่เชื่อมระบบ A กับระบบ B (ใช้คำนวณตารางเชื่อมโยง)',3)
header(wp,4,[('Module ตัวเชื่อม',34),('ระบบ A',16),('ระบบ B',16)])
body(wp,5,[list(p) for p in pairs])
plast=4+len(pairs); page(wp,4,landscape=False)

# ---- Sheet การเชื่อมกัน (matrix + flows)
wm=wb.create_sheet('การเชื่อมกัน')
doms=sorted({d for p in pairs for d in p[1:]})
span=len(doms)+1
title(wm,'แต่ละระบบเชื่อมกันอย่างไร','ตาราง = จำนวน Module ตัวเชื่อม ระหว่าง 2 ระบบ (ยิ่งเยอะ ยิ่งผูกกันแน่น) • นับเฉพาะการเชื่อมตรง (depends ตรง) — ระบบที่มีเงินเกือบทั้งหมดต่อเข้าบัญชีทางอ้อมด้วย ดูคอลัมน์ "ต้องมี App" ในแผ่น ทุก Module',span)
r0=4
c=wm.cell(row=r0,column=1,value='ระบบ'); c.font=f(bold=True,color='FFFFFF'); c.fill=HFILL; c.border=B; c.alignment=CT
wm.column_dimensions['A'].width=16
for j,d in enumerate(doms,2):
    c=wm.cell(row=r0,column=j,value=d); c.font=f(bold=True,color='FFFFFF',size=11); c.fill=HFILL; c.border=B
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True,text_rotation=90)
    wm.column_dimensions[get_column_letter(j)].width=6.5
wm.row_dimensions[r0].height=90
P="'คู่เชื่อม (ข้อมูล)'"
for i,d in enumerate(doms,1):
    rr=r0+i; c=wm.cell(row=rr,column=1,value=d); c.font=f(bold=True); c.border=B; c.fill=PatternFill('solid',fgColor=MID)
    for j in range(2,len(doms)+2):
        col=get_column_letter(j); cc=wm.cell(row=rr,column=j)
        if j-1==i: cc.value='—'
        else:
            cc.value=(f'=COUNTIFS({P}!$B$5:$B${plast},$A{rr},{P}!$C$5:$C${plast},{col}${r0})'
                      f'+COUNTIFS({P}!$B$5:$B${plast},{col}${r0},{P}!$C$5:$C${plast},$A{rr})')
        cc.font=f(); cc.border=B; cc.alignment=Alignment(horizontal='center',vertical='center'); cc.number_format='0;-0;""'
from openpyxl.formatting.rule import ColorScaleRule
mr=f'B{r0+1}:{get_column_letter(len(doms)+1)}{r0+len(doms)}'
wm.conditional_formatting.add(mr, ColorScaleRule(start_type='num',start_value=0,start_color='FFFFFF',end_type='max',end_color='4CAF6E'))
fr=r0+len(doms)+3
wm.cell(row=fr,column=1,value='กระบวนการหลักตั้งแต่ต้นจนจบ (ข้อมูลไหลข้ามระบบอย่างไร)').font=f(size=14,bold=True,color=GREEN)
# flows table placed in its own sheet for print width
page(wm,r0,landscape=True); wm.freeze_panes=None
wm.cell(row=fr+1,column=1,value='ดูแผ่น "กระบวนการหลัก"').font=f(italic=True)

wf=wb.create_sheet('กระบวนการหลัก')
title(wf,'กระบวนการหลัก: ข้อมูลไหลข้ามระบบ + รายการบัญชีที่เกิด','Dr/Cr เป็นภาพรวมตามพฤติกรรมมาตรฐานของ Odoo — ต้องอ่านโค้ดยืนยันอีกครั้งตอนทำ gap รายเรื่อง',4)
header(wf,4,[('กระบวนการ',26),('ขั้นตอน',50),('Module ที่เกี่ยวข้อง',30),('รายการบัญชีที่เกิด (Dr/Cr)',60)])
body(wf,5,[list(x) for x in FLOWS]); page(wf,4)

# ---- Sheet 34 App
wapp=wb.create_sheet('App หลัก')
cols=[('ลำดับ',7),('App',22),('ชื่อเทคนิค',20),('ระบบ',13),('ทำอะไร',38),('ช่วยอะไร',38),('ฟังก์ชันหลัก',70),('ต้องมี App อื่นก่อน',26),('เกี่ยวกับการเงิน',14),('Model ใหม่',10),('เมนู',8),('ตัวเชื่อมที่ใช้ระบบนี้',12)]
title(wapp,f'App หลักของ Odoo Community 20.0',SRC+' • คำอธิบายภาษาไทยสรุปโดยผู้วิเคราะห์',len(cols))
header(wapp,4,cols)
order={'สูง':0,'กลาง':1,'ต่ำ':2}
apps_sorted=sorted(apps,key=lambda r:(order.get(APPS[r['tech']][3].split()[0],3),dom(r['tech']),r['tech']))
data=[]
for i,r in enumerate(apps_sorted,1):
    a=APPS[r['tech']]; rr=4+i
    data.append([i,r['name'],r['tech'],dom(r['tech']),a[0],a[1],a[2],', '.join(r['req_apps']) or '-',a[3],
        f"=INDEX('ทุก Module'!$J$5:$J${last},MATCH(C{rr},'ทุก Module'!$B$5:$B${last},0))",
        f"=INDEX('ทุก Module'!$L$5:$L${last},MATCH(C{rr},'ทุก Module'!$B$5:$B${last},0))",
        f"=COUNTIFS({P}!$B$5:$B${plast},D{rr})+COUNTIFS({P}!$C$5:$C${plast},D{rr})"])
body(wapp,5,data,center_cols=(1,9,10,11,12))
page(wapp,4)

# ---- Community vs Enterprise
we=wb.create_sheet('ไม่มีใน Community')
title(we,'ฟีเจอร์การเงินที่ Community ไม่มี (อยู่ใน Enterprise ซึ่งห้ามใช้อ้างอิง)','ทางเลือก = OCA (odoo-community-association, LGPL/AGPL) — ต้องตรวจว่ามีเวอร์ชันที่ต้องการหรือยัง (OCA มักตามหลัง 1-2 เวอร์ชัน)',4)
header(we,4,[('ฟีเจอร์',50),('Module Enterprise (ห้ามแกะ)',26),('ทางเลือกอ้างอิง (OCA)',45),('ความสำคัญต่อ Laotai',20)])
body(we,5,[list(x) for x in ENTERPRISE],center_cols=(4,)); page(we,4)

# ---- สรุป (first)
ws=wb.create_sheet('สรุป',0)
title(ws,'แผนที่ระบบ Odoo Community 20.0 — สรุป',SRC,4)
ws.column_dimensions['A'].width=34; ws.column_dimensions['B'].width=14; ws.column_dimensions['C'].width=14; ws.column_dimensions['D'].width=60
M="'ทุก Module'"
r=4; header(ws,r,[('กลุ่ม Module (เกณฑ์)',34),('จำนวน',14),('สัดส่วน',14),('ความหมาย',60)])
groups=[('App หลัก','โปรแกรมที่ผู้ใช้เห็นเป็นไอคอนหลัก เช่น บัญชี ขาย คลัง'),
 ('ตัวเชื่อมระหว่างระบบ','Module เล็กๆ ที่ผูก 2 ระบบเข้าหากัน เช่น sale_stock = ขาย↔คลัง'),
 ('ส่วนเสริม','เพิ่มฟีเจอร์ในระบบเดียว'),
 ('ภาษี/ประเทศ (l10n)','ผังบัญชี/ภาษี/รายงานของแต่ละประเทศ (ไม่มีของลาว)'),
 ('เทคนิค/โครงสร้าง','ระบบพื้นฐาน เช่น ผู้ใช้ สิทธิ์ เว็บ อีเมล ไม่ใช่ฟีเจอร์ธุรกิจ'),
 ('โมดูลทดสอบ','ใช้ทดสอบโค้ดเท่านั้น')]
g0=r+1
for i,(g,mean) in enumerate(groups):
    rr=g0+i
    ws.cell(row=rr,column=1,value=g).fill=INFILL
    ws.cell(row=rr,column=2,value=f'=COUNTIF({M}!$D$5:$D${last},A{rr})')
    ws.cell(row=rr,column=3,value=f'=IF($B${g0+len(groups)}=0,0,B{rr}/$B${g0+len(groups)})')
    ws.cell(row=rr,column=4,value=mean)
tr=g0+len(groups)
ws.cell(row=tr,column=1,value='รวมทั้งหมด'); ws.cell(row=tr,column=2,value=f'=SUM(B{g0}:B{tr-1})'); ws.cell(row=tr,column=3,value=f'=SUM(C{g0}:C{tr-1})')
ws.cell(row=tr,column=4,value=f'=IF(B{tr}=COUNTA({M}!$B$5:$B${last}),"ตรวจแล้ว: ตรงกับจำนวนแถวในแผ่น ทุก Module","ไม่ตรง — ตรวจกลุ่ม")')
for rr in range(g0,tr+1):
    for cc in range(1,5):
        c=ws.cell(row=rr,column=cc); c.border=B; c.font=f(bold=(rr==tr)); c.alignment=WR if cc in (1,4) else CT
    ws.cell(row=rr,column=3).number_format='0.0%'
r=tr+2
header(ws,r,[('ระบบ (เกณฑ์)',34),('จำนวน Module',14),('App หลัก',14),('หมายเหตุ',60)])
dlist=sorted({x['domain'] for x in rows}, key=lambda d:-sum(1 for x in rows if x['domain']==d))
d0=r+1
for i,d in enumerate(dlist):
    rr=d0+i
    ws.cell(row=rr,column=1,value=d).fill=INFILL
    ws.cell(row=rr,column=2,value=f'=COUNTIF({M}!$E$5:$E${last},A{rr})')
    ws.cell(row=rr,column=3,value=f'=COUNTIFS({M}!$E$5:$E${last},A{rr},{M}!$D$5:$D${last},"App หลัก")')
    ws.cell(row=rr,column=4,value={'ภาษี/ประเทศ':'แต่ละประเทศ ใช้เป็นตัวอย่างออกแบบผังบัญชี/ภาษีลาว','บัญชี':'กลุ่มหลักที่จะใช้ทำ gap analysis','เทคนิค':'โครงสร้างพื้นฐาน'}.get(d,''))
    for cc in range(1,5):
        c=ws.cell(row=rr,column=cc); c.border=B; c.font=f(); c.alignment=WR if cc in (1,4) else CT
r=d0+len(dlist)+1
notes=['วิธีอ่าน:',
 '• "ฟังก์ชัน" ในไฟล์นี้นับเป็น 2 แบบที่มีความหมายทางธุรกิจ: Model = ตารางข้อมูล/เอกสาร (เช่น ใบแจ้งหนี้), เมนู = หน้าจอที่ผู้ใช้เข้าได้',
 '• ไม่นับจำนวน function ในโค้ด Python เพราะมีเป็นหมื่น และไม่บอกอะไรทางธุรกิจ',
 '• แผ่น "App หลัก" = รายละเอียดภาษาไทยทุก App: ทำอะไร ช่วยอะไร ฟังก์ชันหลัก',
 '• แผ่น "การเชื่อมกัน" + "กระบวนการหลัก" = ระบบไหนผูกกับระบบไหน และข้อมูลไหลอย่างไร',
 '• แผ่น "ไม่มีใน Community" = จุดที่ต้องหาแหล่งอ้างอิงอื่น (OCA)',
 '• เซลล์สีเหลือง = เกณฑ์ที่สูตรใช้ แก้ชื่อกลุ่ม/ระบบได้ ตัวเลขจะคำนวณใหม่']
for i,t in enumerate(notes):
    c=ws.cell(row=r+i,column=1,value=t); c.font=f(bold=(i==0)); ws.merge_cells(start_row=r+i,start_column=1,end_row=r+i,end_column=4)
page(ws,4,landscape=False); ws.freeze_panes=None; ws.print_title_rows=None
wb.save(out); print('saved', len(pairs), 'pairs', len(doms), 'domains')
